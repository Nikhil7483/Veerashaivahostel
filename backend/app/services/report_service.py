import io
import csv
from datetime import datetime
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def generate_pdf_table(title: str, headers: list, data: list, filename: str = "report.pdf", landscape_mode: bool = False):
    buffer = io.BytesIO()
    pagesize = landscape(letter) if landscape_mode else letter
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=6
    )
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=14
    )
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0F172A')
    )
    header_cell_style = ParagraphStyle(
        'TableHeaderCell',
        parent=styles['Normal'],
        fontSize=9,
        leading=11,
        fontName='Helvetica-Bold',
        textColor=colors.white
    )
    
    elements = []
    elements.append(Paragraph(f"Smart Hostel - {title}", title_style))
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Records: {len(data)}", meta_style))
    elements.append(Spacer(1, 10))
    
    table_data = []
    header_row = [Paragraph(str(h), header_cell_style) for h in headers]
    table_data.append(header_row)
    
    for row in data:
        row_cells = [Paragraph(str(cell if cell is not None else "-"), cell_style) for cell in row]
        table_data.append(row_cells)
        
    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3B82F6')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    
    elements.append(t)
    doc.build(elements)
    buffer.seek(0)
    return buffer

def generate_excel_sheet(title: str, headers: list, data: list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:30]
    
    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=f"Smart Hostel System - {title}")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    
    # Subtitle / Timestamp
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sub_cell = ws.cell(row=2, column=1, value=f"Report Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Records: {len(data)}")
    sub_cell.font = Font(name="Calibri", size=9, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20
    
    # Headers
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    ws.row_dimensions[3].height = 24
    for col_idx, header in enumerate(headers, start=1):
        c = ws.cell(row=3, column=col_idx, value=header)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_thin
        
    # Data rows
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    for row_idx, row_values in enumerate(data, start=4):
        ws.row_dimensions[row_idx].height = 20
        use_alt = (row_idx % 2 == 0)
        for col_idx, val in enumerate(row_values, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=str(val if val is not None else "-"))
            c.font = Font(name="Calibri", size=10)
            c.alignment = Alignment(vertical="center")
            c.border = border_thin
            if use_alt:
                c.fill = alt_fill
                
    # Auto adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generate_csv_data(headers: list, data: list):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in data:
        writer.writerow([str(c if c is not None else "") for c in row])
    output.seek(0)
    return output.getvalue()
